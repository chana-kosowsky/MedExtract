"""הערכה מול טבלת האמת. שימוש:
    python eval/evaluate.py run "samples/קובץ.pdf" 1-22 [--runs 3] [--model שם] [--interval שניות]
                                [--media-resolution low|medium|high] [--provider gemini|anthropic|chain]
                                                                     מריץ את המודל, שומר ומחשב ציון
    python eval/evaluate.py score [תיקיית סדרה] [--pages 1-10]        ציון מחדש מתוצאות שמורות, בלי מודל
התוצאות נשמרות ב-eval/results/ (לא נכנס ל-Git), כולל report.xlsx עם רשימת אי-ההתאמות לבדיקה.
"""
import json
import logging
import re
import sys
import threading
import time
from collections import Counter, defaultdict
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "server"))

import openpyxl
import pymupdf
from openpyxl.styles import Font, PatternFill

from app.extractor import (AnthropicExtractor, ChainExtractor, ExtractionError, GeminiExtractor,
                           build_extractor)
from app import pipeline
from app.barcodes import Barcode, read_barcodes, reconcile
from app.pipeline import extract_with_retry, render_pdf_page
from app.schema import ExtractedForm, FieldStatus, ModelField, ModelOutput
from app.validation import build_form, match_sender, parse_date

TRUTH = ROOT / "eval" / "truth_part2.xlsx"
RESULTS = ROOT / "eval" / "results"
TARGET_PAGES, TARGET_SECONDS = 30, 240  # היעד באתגר: 30 טפסים בכ-4 דקות
FIELDS = [
    ("first_name", "שם פרטי"), ("last_name", "שם משפחה"), ("id_number", "ת\"ז"),
    ("central_numbers", "מספרים מרכזיים"), ("received_date", "תאריך קבלה"),
    ("sender", "גורם שולח"), ("commitment_numbers", "מספרי התחייבות"),
]
CORRECT, FALSE_FLAG, CAUGHT, MISSED = "נכון", "סימון שווא", "שגוי וסומן", "שגוי ולא סומן"
CATS = (CORRECT, FALSE_FLAG, CAUGHT, MISSED)


# ---------- טבלת האמת ----------

def parse_pages(spec: str) -> list[int]:
    """ממירה '1-11' או '1,3,5' לרשימת עמודים."""
    pages = []
    for part in spec.split(","):
        a, _, b = part.partition("-")
        pages += range(int(a), int(b or a) + 1)
    return pages


def load_truth() -> dict[int, dict[str, str]]:
    """קוראת את טבלת האמת: עמוד -> {שם עמודה: ערך}."""
    ws = openpyxl.load_workbook(TRUTH)["טבלת אמת"]
    headers = [c.value for c in ws[1]]
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = {h: ("" if v is None else str(v).strip()) for h, v in zip(headers, row)}
        if rec["עמוד"].isdigit():
            rows[int(rec["עמוד"])] = rec
    return rows


def truth_value(attr: str, text: str):
    """מנרמלת ערך אמת לאותה צורה שבה המערכת מחזירה אותו."""
    if attr.endswith("_numbers"):
        return {p.replace(" ", "").replace("-", "") for p in text.split(";") if p.strip()}
    if not text:
        return None
    if attr == "received_date":
        return parse_date(text)[0] or text
    if attr == "sender":
        return match_sender(ModelField(text=text, confidence=1)).value or text
    return text


def difficulty_tags(text: str) -> list[str]:
    """סוגי הקושי של עמוד (מופרדים בפסיק, נקודה-פסיק או לוכסן)."""
    return [t.strip() for t in re.split(r"[,;/]", text or "") if t.strip()]


# ---------- הרצה ----------

def recording(base: type) -> type:
    """עוטפת מחלץ של השרת כך שישמור גם את התשובה הגולמית, כדי שתיקון בכללי הקוד ייבדק בלי לקרוא שוב למודל."""

    class Recording(base):
        def read_page(self, image: bytes, mime_type: str) -> ModelOutput:
            """קוראת עמוד ושומרת את התשובה הגולמית."""
            self.last_raw = super().read_page(image, mime_type)
            return self.last_raw

    return Recording


def run_series(pdf: str, pages: list[int], runs: int, model: str | None = None,
               media_resolution: str | None = None, provider: str = "gemini") -> Path:
    """מריצה את המודל (ברירת מחדל: GEMINI_MODEL / GROQ_MODEL) על העמודים כמה פעמים, כמו בצינור של השרת, ושומרת כל הרצה."""
    if provider == "chain":  # שרשרת המודלים של השרת, לפי MODEL_CHAIN
        built = build_extractor()
        ex = recording(ChainExtractor)(built.links if isinstance(built, ChainExtractor) else [built])
        ex.model = "chain"
    elif provider == "anthropic":
        ex = recording(AnthropicExtractor)(model)
    else:
        ex = recording(GeminiExtractor)(model, media_resolution)
    series = RESULTS / datetime.now().strftime("%Y%m%d-%H%M%S")
    series.mkdir(parents=True)
    never = threading.Event()
    level = getattr(ex, "media_resolution", None)
    resolution = level.name if level else "default"
    print(f"model: {ex.model} | resolution: {resolution} | עמודים: {len(pages)} | הרצות: {runs} | "
          f"תיקייה: {series.name}\n")
    with pymupdf.open(pdf) as doc:
        for n in range(1, runs + 1):
            out = {"model": ex.model, "media_resolution": resolution, "pages": pages, "results": {}}
            started = time.perf_counter()
            for page in pages:
                t0 = time.perf_counter()
                try:
                    image, mime = render_pdf_page(doc, page - 1)
                    form = extract_with_retry(ex, image, mime, page, never)
                    codes = read_barcodes(image)
                    form = reconcile(form, codes)
                    out["results"][page] = {"form": form.model_dump(mode="json"), "model": form.model,
                                            "raw": ex.last_raw.model_dump(mode="json"),
                                            "barcodes": [asdict(b) for b in codes],
                                            "seconds": round(time.perf_counter() - t0, 2)}
                    print(f"  הרצה {n} · עמוד {page}: {out['results'][page]['seconds']} שניות", flush=True)
                except ExtractionError as e:
                    out["results"][page] = {"error": e.code, "message": e.message}
                    print(f"  הרצה {n} · עמוד {page}: נכשל [{e.code}] {e.message}", flush=True)
                    if e.code == "quota":
                        break
            out["wall_seconds"] = round(time.perf_counter() - started, 1)
            (series / f"run_{n}.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
            if any(r.get("error") == "quota" for r in out["results"].values()):
                print("המכסה נגמרה. ההרצות שהושלמו נשמרו.")
                break
    return series


# ---------- ציון ----------

def predicted(form: ExtractedForm, attr: str):
    """מחזירה (ערך, האם סומן, סיבות) של שדה ברשומה."""
    field = getattr(form, attr)
    if isinstance(field, list):
        flagged = not field or any(f.status != FieldStatus.ok for f in field)
        return {f.value for f in field}, flagged, [f.review_reason for f in field if f.review_reason]
    return field.value, field.status != FieldStatus.ok, [field.review_reason] if field.review_reason else []


def classify(truth, pred, flagged) -> str:
    """מסווגת שדה לאחת מארבע הקבוצות."""
    if (truth or None) == (pred or None):
        return FALSE_FLAG if flagged else CORRECT
    return CAUGHT if flagged else MISSED


def show(value) -> str:
    """ערך להצגה: רשימת מספרים ממוינת, או ריק."""
    if isinstance(value, set):
        return "; ".join(sorted(v or "?" for v in value)) or "—"
    return value or "—"


def load_series(series: Path) -> list[dict]:
    """טוענת את כל ההרצות השמורות בסדרה."""
    runs = []
    for f in sorted(series.glob("run_*.json")):
        data = json.loads(f.read_text(encoding="utf-8"))
        data["results"] = {int(k): v for k, v in data["results"].items()}
        runs.append(data)
    return runs


def score(series: Path, only: list[int] | None = None) -> None:
    """מחשבת ציון לכל ההרצות בסדרה מול טבלת האמת (או רק לעמודים שנבחרו), מדפיסה סיכום ושומרת report.xlsx."""
    runs, truth = load_series(series), load_truth()
    if not runs:
        sys.exit(f"אין הרצות שמורות ב-{series}")
    if only:
        truth = {p: t for p, t in truth.items() if p in only}
    print(f"model: {runs[0]['model']} | resolution: {runs[0].get('media_resolution', 'default')} | "
          f"הרצות: {len(runs)} | סדרה: {series.name}\n")

    counts = {attr: Counter() for attr, _ in FIELDS}
    per_run = [Counter() for _ in runs]
    by_tag = defaultdict(Counter)
    cells = defaultdict(list)  # (עמוד, שדה) -> [(ערך, סומן, קטגוריה, סיבות) לכל הרצה]
    seconds, failed = [], Counter()
    for i, run in enumerate(runs):
        for page, res in sorted(run["results"].items()):
            if "error" in res:
                failed[res["error"]] += 1
                continue
            if page not in truth:
                continue
            seconds.append(res["seconds"])
            # אם נשמרה התשובה הגולמית, הכללים בקוד מופעלים מחדש (כך נמדד תיקון בקוד בלי מודל)
            if "raw" in res:
                form = build_form(ModelOutput.model_validate(res["raw"]), page)
                form = reconcile(form, [Barcode(**b) for b in res.get("barcodes", [])])
            else:
                form = ExtractedForm.model_validate(res["form"])
            tags = difficulty_tags(truth[page].get("סוגי קושי", "")) or ["ללא"]
            for attr, label in FIELDS:
                t = truth_value(attr, truth[page][label])
                p, flagged, reasons = predicted(form, attr)
                cat = classify(t, p, flagged)
                counts[attr][cat] += 1
                per_run[i][cat] += 1
                for tag in tags:
                    by_tag[tag][cat] += 1
                cells[(page, label)].append((p, flagged, cat, reasons))

    total = sum(counts.values(), Counter())
    print("=== סיכום לפי שדה (כל ההרצות יחד) ===")
    print(f"{'שדה':16}| {CORRECT:6}| {FALSE_FLAG:10}| {CAUGHT:10}| {MISSED}")
    for attr, label in FIELDS:
        c = counts[attr]
        print(f"{label:16}| {c[CORRECT]:6}| {c[FALSE_FLAG]:10}| {c[CAUGHT]:10}| {c[MISSED]}")
    print(f"{'סה\"כ':16}| {total[CORRECT]:6}| {total[FALSE_FLAG]:10}| {total[CAUGHT]:10}| {total[MISSED]}")

    print("\n=== לפי הרצה ===")
    for i, c in enumerate(per_run, 1):
        print(f"  הרצה {i}: " + " | ".join(f"{k} {c[k]}" for k in CATS))

    print("\n=== לפי סוג קושי ===")
    for tag, c in sorted(by_tag.items(), key=lambda kv: -sum(kv[1].values())):
        n = sum(c.values())
        print(f"  {tag}: {n} שדות | נכון (עם או בלי סימון) {c[CORRECT] + c[FALSE_FLAG]} | "
              f"שגוי ולא סומן {c[MISSED]}")

    unstable = [(k, v) for k, v in cells.items() if len({show(x[0]) for x in v}) > 1]
    mismatches = [(k, v) for k, v in sorted(cells.items()) if any(x[2] in (CAUGHT, MISSED) for x in v)]
    print(f"\n=== יציבות: {len(unstable)} שדות קיבלו ערך שונה בין ההרצות (מתוך {len(cells)}) ===")
    for (page, label), v in sorted(unstable):
        print(f"  עמוד {page} | {label} | " + " / ".join(show(x[0]) for x in v))

    attr_of = {label: attr for attr, label in FIELDS}
    print(f"\n=== אי-התאמות לטבלת האמת: {len(mismatches)} (שגוי ולא סומן מסומן ב-!) ===")
    for (page, label), v in mismatches:
        mark = "!" if any(x[2] == MISSED for x in v) else " "
        t = show(truth_value(attr_of[label], truth[page][label]))
        print(f" {mark} עמוד {page} | {label} | אמת: {t} | מערכת: " + " / ".join(show(x[0]) for x in v))

    if seconds:
        avg = sum(seconds) / len(seconds)
        walls = [r["wall_seconds"] / len(r["pages"]) for r in runs if r.get("wall_seconds")]
        per_page = sum(walls) / len(walls) if walls else avg
        print(f"\nזמן ממוצע לעמוד: {avg:.1f} שניות במודל, {per_page:.1f} שניות כולל המתנות | "
              f"{TARGET_PAGES} עמודים ≈ {per_page * TARGET_PAGES / 60:.1f} דקות (יעד: {TARGET_SECONDS // 60})")
    if failed:
        print("עמודים שנכשלו: " + ", ".join(f"{k} {v}" for k, v in failed.items()))

    if not only:  # ציון חלקי לא דורס את הדוח המלא
        path = write_report(series, runs, truth, counts, mismatches, unstable)
        print(f"\nדוח לבדיקה: {path}")


def _cell(ws, row: int, col: int, value, fill=None) -> None:
    """כותבת תא כטקסט בלבד (בלי נוסחאות ובלי לאבד אפס מוביל)."""
    c = ws.cell(row=row, column=col, value=value)
    if value is not None:
        c.data_type = "s"
    if fill:
        c.fill = fill


def write_report(series: Path, runs, truth, counts, mismatches, unstable) -> Path:
    """קובץ אקסל לבדיקת המפתחת: סיכום, אי-התאמות מול הטבלה (עם עמודה לסימון), ויציבות."""
    attr_of = {label: attr for attr, label in FIELDS}
    red, orange = PatternFill("solid", fgColor="F8C9C9"), PatternFill("solid", fgColor="FFE3C2")
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "סיכום"
    ws.sheet_view.rightToLeft = True
    for col, title in enumerate(["שדה", *CATS], 1):
        ws.cell(1, col, title).font = Font(bold=True)
    for r, (attr, label) in enumerate(FIELDS, 2):
        ws.cell(r, 1, label)
        for col, cat in enumerate(CATS, 2):
            ws.cell(r, col, counts[attr][cat])

    ws = wb.create_sheet("אי-התאמות")
    ws.sheet_view.rightToLeft = True
    head = ["עמוד", "שדה", "בטבלת האמת"] + [f"הרצה {i}" for i in range(1, len(runs) + 1)] + \
           ["סומן לבדיקה?", "סיבה שהמערכת נתנה", "מי צודק? (טבלה / מערכת)"]
    for col, title in enumerate(head, 1):
        ws.cell(1, col, title).font = Font(bold=True)
    for r, ((page, label), v) in enumerate(mismatches, 2):
        _cell(ws, r, 1, str(page))
        _cell(ws, r, 2, label)
        _cell(ws, r, 3, show(truth_value(attr_of[label], truth[page][label])))
        for i, (p, flagged, cat, _) in enumerate(v):
            _cell(ws, r, 4 + i, show(p), red if cat == MISSED else orange if cat == CAUGHT else None)
        _cell(ws, r, 4 + len(runs), " / ".join("כן" if x[1] else "לא" for x in v))
        _cell(ws, r, 5 + len(runs), " | ".join(sorted({"; ".join(x[3]) for x in v if x[3]})))
    for col, width in zip("ABC", (7, 16, 22)):
        ws.column_dimensions[col].width = width

    ws = wb.create_sheet("יציבות")
    ws.sheet_view.rightToLeft = True
    for col, title in enumerate(["עמוד", "שדה"] + [f"הרצה {i}" for i in range(1, len(runs) + 1)], 1):
        ws.cell(1, col, title).font = Font(bold=True)
    for r, ((page, label), v) in enumerate(sorted(unstable), 2):
        _cell(ws, r, 1, str(page))
        _cell(ws, r, 2, label)
        for i, x in enumerate(v):
            _cell(ws, r, 3 + i, show(x[0]))

    path = series / "report.xlsx"
    wb.save(path)
    return path


def main() -> None:
    """run: הרצה ושמירה ואז ציון; score: ציון מחדש של סדרה שמורה (ברירת מחדל: האחרונה)."""
    sys.stdout.reconfigure(encoding="utf-8")
    logging.basicConfig(level=logging.INFO, format="  [log] %(message)s")  # ניסיונות חוזרים, בלי תוכן
    args = sys.argv[1:]
    if args[:1] == ["run"] and len(args) >= 3:
        runs = int(args[args.index("--runs") + 1]) if "--runs" in args else 3
        model = args[args.index("--model") + 1] if "--model" in args else None
        resolution = args[args.index("--media-resolution") + 1] if "--media-resolution" in args else None
        provider = args[args.index("--provider") + 1] if "--provider" in args else "gemini"
        if "--interval" in args:  # מרווח בין בקשות לפי מכסת הדקה של המודל (Flash חינמי: 5 בדקה)
            pipeline.MIN_INTERVAL = float(args[args.index("--interval") + 1])
        score(run_series(args[1], parse_pages(args[2]), runs, model, resolution, provider))
    elif args[:1] == ["score"]:
        only = parse_pages(args.pop(args.index("--pages") + 1)) if "--pages" in args else None
        args = [a for a in args if a != "--pages"]
        series = Path(args[1]) if len(args) > 1 else max(RESULTS.iterdir(), key=lambda p: p.name)
        score(series, only)
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main()
