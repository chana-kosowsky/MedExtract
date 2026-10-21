from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session, select

from app import db
from app.main import app
from app.schema import Commitment
from tests.test_db import REVIEW, fv, make_form


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture(autouse=True)
def clean():
    """כל בדיקה מתחילה מטבלה ריקה (בסכמת הבדיקה בלבד)."""
    with Session(db.get_engine()) as s:
        db._delete_ids(s, list(s.exec(select(Commitment.id)).all()))
        s.commit()


def _book(client, **params):
    """מורידה את הייצוא ומחזירה את הגליונות."""
    res = client.get("/api/commitments/export", params=params)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in res.headers["content-disposition"]
    wb = load_workbook(BytesIO(res.content))
    return wb["התחייבויות"], wb["פירוט"]


def test_export_rows_columns_and_detail(client):
    """שורה לכל התחייבות לפי סדר האתגר, מספרים מופרדים בפסיק, ושורה לכל מספר בפירוט."""
    db.save_form(make_form(), "b1", 1)
    main, detail = _book(client)
    header = [c.value for c in main[1]]
    assert header[:8] == ["מצב", "שם פרטי", "שם משפחה", 'ת"ז', "מספר מרכזי", "תאריך קבלה",
                          "גורם שולח", "מספר התחייבות"]
    assert header[8:] == ["מקור", "נוצר", "אושר"]
    assert "הערות לבדיקה" not in header
    row = [c.value for c in main[2]]
    assert row[1:8] == ["ישראל", "כהן", "011223344", "947468, 947469", "13/08/2026", "שניידר", "12345678"]
    assert main.sheet_view.rightToLeft
    assert detail.max_row == 1 + 3


def test_review_cells_orange_only_until_approved(client):
    """תא לבדיקה בכתום ותא תקין בלי צבע; אחרי אישור הרשומה נקייה. אין הערות בקובץ."""
    rid = db.save_form(make_form(id_number=fv("011223345", REVIEW, "ספרת ביקורת לא תקינה")), "b1", 1)
    main, _ = _book(client)
    assert main.cell(2, 4).fill.fgColor.rgb.endswith("FFD8A8")
    assert main.cell(2, 6).fill.fill_type is None
    assert main.max_column == 11
    db.approve(rid)
    main, _ = _book(client)
    assert all(c.fill.fill_type is None for c in main[2])


def test_values_stay_text(client):
    """אפס מוביל נשמר, וערך שנראה כמו נוסחה נכתב כטקסט ולא מורץ."""
    db.save_form(make_form(first_name=fv("=HYPERLINK(\"x\")", REVIEW, "שם")), "b1", 1)
    main, _ = _book(client)
    assert main.cell(2, 4).value == "011223344"
    assert main.cell(2, 2).data_type == "s" and main.cell(2, 2).value.startswith("=")


def test_export_filters_and_skips_failed(client):
    """הייצוא מכבד את הסינון, ועמוד שנכשל לא נכלל."""
    db.save_form(make_form(), "b1", 1)
    db.save_form(make_form(sender=fv("וולפסון")), "b2", 1)
    db.save_failure("b1", 1, 2, "bad_page", "העמוד פגום")
    main, _ = _book(client)
    assert main.max_row == 1 + 2
    main, _ = _book(client, batch_id="b2")
    assert main.max_row == 1 + 1 and main.cell(2, 7).value == "וולפסון"


def test_export_empty(client):
    """בלי רשומות: קובץ עם כותרות בלבד."""
    main, detail = _book(client)
    assert main.max_row == 1 and detail.max_row == 1
