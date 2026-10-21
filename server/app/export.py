from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

# סדר השדות של האתגר (החלטה 9)
COLUMNS = [("first_name", "שם פרטי"), ("last_name", "שם משפחה"), ("id_number", 'ת"ז'),
           ("central", "מספר מרכזי"), ("received_date", "תאריך קבלה"), ("sender", "גורם שולח"),
           ("commitment", "מספר התחייבות")]
NUMBER_LABELS = {"central": "מספר מרכזי", "commitment": "מספר התחייבות"}
STATE_LABELS = {"pending": "ממתין", "approved": "אושר", "failed": "נכשל"}

REVIEW_FILL = PatternFill("solid", fgColor="FFD8A8")  # תא לבדיקה בכתום (החלטה 17)
HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LOCAL_TZ = ZoneInfo("Asia/Jerusalem")


def _when(dt) -> str:
    """זמן בשעון ישראל, או ריק."""
    return dt.astimezone(LOCAL_TZ).strftime("%d/%m/%Y %H:%M") if dt else ""


def _write(ws: Worksheet, row: int, col: int, value, review: bool = False) -> None:
    """כותבת תא כטקסט בלבד (בלי נוסחאות, ובלי לאבד אפס מוביל), ומסמנת בכתום אם צריך."""
    cell = ws.cell(row=row, column=col, value=value if value not in (None, "") else None)
    if cell.value is not None:
        cell.data_type = "s"
    cell.number_format = "@"
    if review:
        cell.fill = REVIEW_FILL


def _header(ws: Worksheet, titles: list[str], widths: list[int]) -> None:
    """שורת כותרת מעוצבת, מוקפאת, מימין לשמאל."""
    ws.sheet_view.rightToLeft = True
    for i, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def to_xlsx(items: list[dict]) -> bytes:
    """קובץ אקסל נקי: התחייבויות ופירוט מספרים, כתום רק במה שעוד לא אושר, בלי הערות (החלטה 17)."""
    wb = Workbook()
    main = wb.active
    main.title = "התחייבויות"
    titles = ["מצב"] + [t for _, t in COLUMNS] + ["מקור", "נוצר", "אושר"]
    _header(main, titles, [9, 14, 16, 12, 22, 13, 22, 22, 16, 17, 17])

    detail = wb.create_sheet("פירוט")
    _header(detail, ["שם פרטי", "שם משפחה", 'ת"ז', "סוג", "מספר"], [14, 16, 12, 16, 16])

    drow = 2
    for r, item in enumerate(items, start=2):
        pending = item["state"] != "approved"
        _write(main, r, 1, STATE_LABELS.get(item["state"], item["state"]))
        for c, (key, _) in enumerate(COLUMNS, start=2):
            if key in NUMBER_LABELS:
                numbers = item[f"{key}_numbers"]
                flagged = not numbers or any(n["status"] != "ok" for n in numbers)
                _write(main, r, c, ", ".join(n["value"] or "?" for n in numbers), pending and flagged)
            else:
                field = item[key]
                _write(main, r, c, field["value"], pending and field["status"] != "ok")
        page = item.get("source_page")
        source = "דוגמה" if item["kind"] == "sample" else (f"עמוד {page}" if page else "")
        _write(main, r, 9, source)
        _write(main, r, 10, _when(item.get("created_at")))
        _write(main, r, 11, _when(item.get("approved_at")))

        for key, label in NUMBER_LABELS.items():
            for n in item[f"{key}_numbers"]:
                for c, v in enumerate((item["first_name"]["value"], item["last_name"]["value"],
                                       item["id_number"]["value"], label), start=1):
                    _write(detail, drow, c, v)
                _write(detail, drow, 5, n["value"], pending and n["status"] != "ok")
                drow += 1

    main.auto_filter.ref = main.dimensions
    detail.auto_filter.ref = detail.dimensions
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
